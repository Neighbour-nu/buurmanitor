"""Een artikel uit de cockpit als Word-document (.docx), bijvoorbeeld om een persbericht te mailen of te bewaren."""
import io
import re

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import opmaak

GROEN = RGBColor(0x2E, 0x9E, 0x45)
INLINE = re.compile(r"\[([^\]]+)\]\((https?://[^\s)]+)\)|\*\*(.+?)\*\*|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])")


def _link(alinea, tekst, url):
    rid = alinea.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    kleur = OxmlElement("w:color")
    kleur.set(qn("w:val"), "1F4E9A")
    onder = OxmlElement("w:u")
    onder.set(qn("w:val"), "single")
    rpr.append(kleur)
    rpr.append(onder)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = tekst
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    alinea._p.append(h)


def _inline(alinea, tekst):
    pos = 0
    for m in INLINE.finditer(tekst):
        if m.start() > pos:
            alinea.add_run(tekst[pos:m.start()])
        if m.group(1):
            _link(alinea, m.group(1), m.group(2))
        elif m.group(3):
            alinea.add_run(m.group(3)).bold = True
        else:
            alinea.add_run(m.group(4)).italic = True
        pos = m.end()
    if pos < len(tekst):
        alinea.add_run(tekst[pos:])


def maak(titel, tekst, label=None):
    doc = Document()
    stijl = doc.styles["Normal"]
    stijl.font.name = "Calibri"
    stijl.font.size = Pt(11)
    for s in ("Title", "Heading 2"):
        doc.styles[s].font.name = "Calibri"
    if label:
        p = doc.add_paragraph()
        r = p.add_run(label.upper())
        r.bold, r.font.size, r.font.color.rgb = True, Pt(10), GROEN
    doc.add_heading(titel or "Zonder titel", level=0)
    for soort, inhoud in opmaak.blokken(tekst or ""):
        if soort == "kop":
            doc.add_heading(inhoud, level=2)
        elif soort == "p":
            _inline(doc.add_paragraph(), inhoud)
        elif soort in ("ul", "ol"):
            for x in inhoud:
                _inline(doc.add_paragraph(style="List Bullet" if soort == "ul" else "List Number"), x)
        elif soort == "tabel":
            kolommen = max(len(r) for r in inhoud)
            t = doc.add_table(rows=len(inhoud), cols=kolommen)
            t.style = "Table Grid"
            for i, rij in enumerate(inhoud):
                for j, cel in enumerate(rij):
                    alinea = t.cell(i, j).paragraphs[0]
                    _inline(alinea, cel)
                    if i == 0:
                        for r in alinea.runs:
                            r.bold = True
        elif soort == "plek":
            p = doc.add_paragraph()
            r = p.add_run(f"[Vast onderdeel van de website: {inhoud[1]}]")
            r.italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def bestandsnaam(titel):
    naam = re.sub(r"[^\w\s-]", "", (titel or "artikel")).strip()
    naam = re.sub(r"\s+", " ", naam)[:80] or "artikel"
    return naam + ".docx"
