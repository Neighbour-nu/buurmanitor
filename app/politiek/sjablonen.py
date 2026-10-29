"""De formats van de griffie invullen: amendement, motie, technische vraag, politieke vraag
en vragen bij een geagendeerd onderwerp.

De sjablonen zijn de Word-bestanden van de griffie. Promonitor vult ze in en laat opmaak,
logo's en de noten van de griffie staan. Een nieuwe versie van de griffie upload je op het
scherm Sjablonen; die gaat dan voor op de standaardversie in app/sjablonen/.

Elke soort krijgt een dict met gegevens (zie VELDEN). Het bijdragenatelier (fase 1) levert
precies zo'n dict aan, zodat elk amendement, elke motie en elke vraag in het juiste format komt.
"""
import copy
import io
import os
import re

import docx
from docx.shared import Cm
from docx.text.paragraph import Paragraph

MAP = os.path.join(os.path.dirname(__file__), "sjablonen")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SOORTEN = {
    "amendement": "Amendement",
    "motie": "Motie",
    "technische_vraag": "Technische vragen",
    "politieke_vraag": "Politieke vragen (schriftelijke vragen aan het college)",
    "vraag_agendapunt": "Vragen bij een geagendeerd onderwerp",
}

# Velden voor het invulformulier. Soort veld: tekst (één regel), regels (één punt per regel), keuze.
VELDEN = {
    "amendement": [
        ("raadsvoorstel_titel", "Titel van het raadsvoorstel", "tekst"),
        ("raadsvoorstel_nummer", "Nummer van het raadsvoorstel (bijvoorbeeld 2026-069)", "tekst"),
        ("vergaderdatum", "Datum van de raadsvergadering (bijvoorbeeld 27 oktober 2026)", "tekst"),
        ("constateren", "Wat we constateren (één punt per regel)", "regels"),
        ("veranderen", "Wat willen we veranderen en waarom (één punt per regel)", "regels"),
        ("wijziging", "Soort wijziging", ("keuze", [("veranderen", "Verandert het besluit (van ... in ...)"),
                                                   ("toevoegen", "Voegt beslispunten toe"),
                                                   ("verwijderen", "Verwijdert beslispunten")])),
        ("van", "Bij veranderen: de huidige tekst van het besluit", "regels"),
        ("in", "Bij veranderen: de nieuwe tekst (één alinea of onderdeel per regel)", "regels"),
        ("punten", "Bij toevoegen of verwijderen: de beslispunten (één per regel)", "regels"),
        ("indieners", "Indieners (één per regel; bij een andere partij: naam | partij)", "regels"),
    ],
    "motie": [
        ("titel", "Onderwerp in de kop (bijvoorbeeld Betaalbaar sporten; mag leeg)", "tekst"),
        ("raadsvoorstel_titel", "Titel van het raadsvoorstel (leeg bij een motie vreemd aan de orde)", "tekst"),
        ("raadsvoorstel_nummer", "Nummer van het raadsvoorstel", "tekst"),
        ("vreemd_onderwerp", "Of: onderwerp dat niet op de agenda staat (motie vreemd aan de orde)", "tekst"),
        ("vergaderdatum", "Datum van de raadsvergadering", "tekst"),
        ("vaststellen", "We stellen vast dat (één punt per regel)", "regels"),
        ("overwegingen", "We hebben nagedacht over en houden rekening met (één punt per regel)", "regels"),
        ("vinden", "We vinden dat (één punt per regel)", "regels"),
        ("opdracht_soort", "Soort opdracht", ("keuze", [("vragen", "We vragen het college om:"),
                                                       ("uitspreken", "We spreken uit dat:"),
                                                       ("opdracht", "We geven het college opdracht om:")])),
        ("dictum", "De opdracht of uitspraak (één punt per regel)", "regels"),
        ("indieners", "Indieners (één per regel; bij een andere partij: naam | partij)", "regels"),
    ],
    "technische_vraag": [
        ("onderwerp", "Onderwerp (bijvoorbeeld Raadsvoorstel 2026-067, Bouwplan Baroniehof)", "tekst"),
        ("portefeuillehouder", "Portefeuillehouder (mag leeg)", "tekst"),
        ("vragen", "Vragen (één per regel; ## kop voor een tussenkop, - voor een deelvraag)", "regels"),
        ("indieners", "Indieners (één per regel)", "regels"),
    ],
    "politieke_vraag": [
        ("onderwerp", "Onderwerp", "tekst"),
        ("inleiding", "Inleiding (één alinea per regel)", "regels"),
        ("vragen", "Vragen (één per regel; - voor een deelvraag)", "regels"),
        ("indieners", "Indieners (één per regel)", "regels"),
    ],
    "vraag_agendapunt": [
        ("vergadering", "Vergadering (bijvoorbeeld Commissie Ruimte 6 oktober 2026)", "tekst"),
        ("agendapunt", "Agendapunt (bijvoorbeeld 5. Raadsvoorstel 2026-067, Bouwplan Baroniehof)", "tekst"),
        ("portefeuillehouder", "Portefeuillehouder (mag leeg)", "tekst"),
        ("vragen", "Vragen (één per regel; - voor een deelvraag)", "regels"),
        ("indieners", "Indieners (één per regel)", "regels"),
    ],
}

VOORBEELD = {
    "amendement": {
        "raadsvoorstel_titel": "Afwegingskader sportaccommodaties", "raadsvoorstel_nummer": "2026-069",
        "vergaderdatum": "27 oktober 2026",
        "constateren": ["Clubs betalen na deprivatisering huur aan de gemeente.",
                        "Het kader kijkt niet naar de gevolgen voor de contributie."],
        "veranderen": ["We leggen vast dat de nieuwe huur past binnen de begroting van de club."],
        "wijziging": "veranderen",
        "van": ["\"Het afwegingskader inclusief bijlage vast te stellen.\""],
        "in": ["\"Het afwegingskader inclusief bijlage vast te stellen, met de volgende wijziging:",
               "a. Bijlage 0, uitgangspunt 8: na 'betaalbare lasten voor de sportvereniging' toevoegen 'en haar leden'.\""],
        "indieners": ["Martijn Buurman"],
    },
    "motie": {
        "titel": "", "raadsvoorstel_titel": "Afwegingskader sportaccommodaties", "raadsvoorstel_nummer": "2026-069",
        "vergaderdatum": "27 oktober 2026",
        "vaststellen": ["Sporten moet voor iedere inwoner betaalbaar blijven."],
        "overwegingen": ["De Sport- en beweegnota 2021-2030 wil drempels wegnemen."],
        "vinden": ["Geen kind mag buitenspel staan vanwege geldgebrek."],
        "opdracht_soort": "vragen",
        "dictum": ["Jaarlijks te rapporteren over de contributie bij gedeprivatiseerde clubs."],
        "indieners": ["Martijn Buurman"],
    },
    "technische_vraag": {
        "onderwerp": "Raadsvoorstel 2026-067, Bouwplan Baroniehof", "portefeuillehouder": "",
        "vragen": ["## Bomen", "Hoeveel bomen worden er in totaal gekapt?",
                   "Is onderzocht of bomen verplant kunnen worden?", "- Zo ja, welke?", "- Zo nee, waarom niet?",
                   "## Overige onderwerpen", "Welke geluidsmaatregel is gekozen?"],
        "indieners": ["Martijn Buurman"],
    },
    "politieke_vraag": {
        "onderwerp": "Digitale weerbaarheid",
        "inleiding": ["Gemeenten worden steeds vaker getroffen door storingen en aanvallen.",
                      "Wij willen weten hoe Helmond hierop is voorbereid."],
        "vragen": ["Is er een actueel continuïteitsplan?", "- Zo ja, wanneer is het voor het laatst geoefend?",
                   "- Zo nee, wanneer komt het er?", "Hoe wordt de raad geïnformeerd bij een grote storing?"],
        "indieners": ["Martijn Buurman"],
    },
    "vraag_agendapunt": {
        "vergadering": "Commissie, 6 oktober 2026", "agendapunt": "Raadsvoorstel 2026-067, Bouwplan Baroniehof",
        "portefeuillehouder": "",
        "vragen": ["Hoeveel bomen worden er gekapt?", "Hoeveel nieuwe bomen komen ervoor terug?"],
        "indieners": ["Martijn Buurman"],
    },
}

STANDAARD_FRACTIE = "PRO Helmond"


# ---------- algemene hulpfuncties voor Word ----------

def _plat_velden(doc):
    """Vervangt de invulvelden van Word (FORMTEXT) door gewone tekst, zodat we die kunnen vervangen."""
    for p in doc.element.body.iter(W + "p"):
        kinderen = list(p)
        i = 0
        while i < len(kinderen):
            r = kinderen[i]
            fc = r.find(W + "fldChar") if r.tag == W + "r" else None
            if fc is None or fc.get(W + "fldCharType") != "begin":
                i += 1
                continue
            # Zoek separate en end bij dit veld.
            j, scheiding, einde = i + 1, None, None
            while j < len(kinderen):
                f = kinderen[j].find(W + "fldChar") if kinderen[j].tag == W + "r" else None
                if f is not None and f.get(W + "fldCharType") == "separate":
                    scheiding = j
                if f is not None and f.get(W + "fldCharType") == "end":
                    einde = j
                    break
                j += 1
            instr = "".join(t.text or "" for k in kinderen[i:scheiding or j] for t in k.iter(W + "instrText"))
            if einde is None or "FORMTEXT" not in instr:
                i += 1
                continue
            for k in kinderen[i:(scheiding if scheiding is not None else einde) + 1] + [kinderen[einde]]:
                if k.tag == W + "r":
                    p.remove(k)
            kinderen = list(p)
            i = 0


def _alinea(doc, begin):
    """De eerste alinea waarvan de tekst (zonder spaties en tabs) begint met `begin`."""
    for p in doc.paragraphs:
        if p.text.strip().startswith(begin):
            return p
    raise ValueError(f"Alinea '{begin}' niet gevonden in het sjabloon")


def _index(doc, p):
    return next(i for i, x in enumerate(doc.paragraphs) if x._p is p._p)


def _zet(p, tekst, cursief=None, vet=None):
    """Zet de tekst van een alinea; de opmaak van de eerste tekstrun blijft."""
    runs = [r for r in p.runs if not r._r.xpath(".//w:drawing|.//w:pict")]
    if not runs:
        r = p.add_run(tekst)
    else:
        r = next((x for x in runs if x.text.strip()), runs[0])
        for x in runs:
            if x is not r:
                x._r.getparent().remove(x._r)
        r.text = tekst
    if cursief is not None:
        r.italic = cursief
    if vet is not None:
        r.bold = vet
    return p


def _vervang(p, oud, nieuw):
    """Vervangt de tekst van een run die precies `oud` is (een voormalig invulveld)."""
    for r in p.runs:
        if r.text == oud:
            r.text = nieuw
            return True
    return False


def _veld(p, waarde):
    """Voor regels als 'Onderwerp<tab>: ...': alles na de dubbele punt wordt `waarde`."""
    runs = p.runs
    for i, r in enumerate(runs):
        if ":" in r.text:
            r.text = r.text[: r.text.index(":") + 1] + " " + (waarde or "")
            for x in runs[i + 1:]:
                x._r.getparent().remove(x._r)
            return
    raise ValueError(f"Geen veld in alinea '{p.text}'")


def _kloon_na(p, prototype):
    """Zet een kopie van alinea `prototype` direct na `p` en geeft de nieuwe alinea terug."""
    nieuw = copy.deepcopy(prototype._p)
    p._p.addnext(nieuw)
    return Paragraph(nieuw, p._parent)


def _weg(p):
    p._p.getparent().remove(p._p)


def _leeg(p):
    return not p.text.strip() and not p._p.xpath(".//w:drawing|.//w:pict")


def _lijst(p, items):
    """Vult een (genummerde) placeholder-alinea met items; extra items als kopie eronder."""
    items = [i for i in (items or []) if str(i).strip()] or [""]
    proto = copy.deepcopy(p._p)
    _zet(p, items[0])
    laatste = p
    for item in items[1:]:
        nieuw = copy.deepcopy(proto)
        laatste._p.addnext(nieuw)
        laatste = _zet(Paragraph(nieuw, p._parent), item)
    return laatste


def _ruim_lege_lijstpunten(doc):
    """Haalt lege genummerde alinea's weg (die zouden een leeg opsommingsteken tonen)."""
    for p in list(doc.paragraphs):
        if p._p.pPr is not None and p._p.pPr.numPr is not None and _leeg(p):
            _weg(p)


def _indieners(gegevens):
    uit = []
    for regel in gegevens.get("indieners") or ["Martijn Buurman"]:
        naam, _, fractie = str(regel).partition("|")
        if naam.strip():
            uit.append((naam.strip(), fractie.strip() or STANDAARD_FRACTIE))
    return uit or [("Martijn Buurman", STANDAARD_FRACTIE)]


def _zet_indieners(doc, gegevens):
    """De regels 'Naam' en 'Fractie' onder Indiener(s) in amendement en motie."""
    ind = _indieners(gegevens)
    naam_p, fractie_p = _alinea(doc, "Naam"), _alinea(doc, "Fractie")
    fracties = list(dict.fromkeys(f for _, f in ind))
    if len(fracties) == 1:
        _lijst(naam_p, [n for n, _ in ind])
        _zet(fractie_p, fracties[0])
    else:
        _lijst(naam_p, [f"{n} ({f})" for n, f in ind])
        _weg(fractie_p)


def _namen(gegevens):
    ind = _indieners(gegevens)
    namen = [n for n, _ in ind]
    tekst = namen[0] if len(namen) == 1 else ", ".join(namen[:-1]) + " en " + namen[-1]
    return tekst, ", ".join(dict.fromkeys(f for _, f in ind))


def _vragen_structuur(regels):
    """Regels naar een lijst van ('kop', tekst), ('vraag', tekst) en ('sub', tekst)."""
    uit = []
    for regel in regels or []:
        r = str(regel).rstrip()
        if not r.strip():
            continue
        if r.lstrip().startswith("##"):
            uit.append(("kop", r.lstrip("# ").strip()))
        elif r.lstrip().startswith(("- ", "* ", "• ")) or r.startswith(("  ", "\t")):
            uit.append(("sub", r.lstrip(" \t-*•").strip()))
        else:
            uit.append(("vraag", re.sub(r"^\d+[./)]\s*", "", r.strip())))
    return uit


def _schrijf_vragen(start, proto_leeg, regels, stijl="punt"):
    """Zet vragen na alinea `start`: genummerd, met tussenkoppen en deelvragen (a, b, c).
    stijl 'punt' geeft '1.', stijl 'schuin' geeft '1/' zoals in de politieke vragen van de griffie."""
    laatste, nr, sub = start, 0, 0
    for soort, tekst in _vragen_structuur(regels):
        if soort == "kop":
            laatste = _kloon_na(laatste, proto_leeg)
            laatste = _zet(_kloon_na(laatste, proto_leeg), tekst, vet=True, cursief=False)
            continue
        if soort == "vraag":
            nr, sub = nr + 1, 0
            laatste = _kloon_na(laatste, proto_leeg)
            label = f"{nr}." if stijl == "punt" else f"{nr}/"
            p = _zet(_kloon_na(laatste, proto_leeg), f"{label}\t{tekst}", vet=False, cursief=False)
            inspring = Cm(0.75)
        else:
            sub += 1
            p = _zet(_kloon_na(laatste, proto_leeg), f"{chr(96 + sub)}.\t{tekst}", vet=False, cursief=False)
            inspring = Cm(1.5)
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = inspring, -Cm(0.75)
        pf.tab_stops.add_tab_stop(inspring)
        laatste = p
    return laatste


def _leeg_tussen(doc, start, eind, houd=1):
    """Haalt de lege alinea's tussen `start` en `eind` weg; houdt er `houd` vlak voor `eind`."""
    a, b = _index(doc, start), _index(doc, eind)
    tussen = [p for p in doc.paragraphs[a + 1:b] if _leeg(p)]
    for p in tussen[: max(0, len(tussen) - houd)]:
        _weg(p)


# ---------- per soort ----------

def _amendement(doc, g):
    p = _alinea(doc, "Behorend bij raadsvoorstel")
    _vervang(p, "titel agendapunt", g.get("raadsvoorstel_titel", ""))
    _vervang(p, "nummer", g.get("raadsvoorstel_nummer", ""))
    _vervang(_alinea(doc, "De gemeenteraad van Helmond"), "datum", g.get("vergaderdatum", ""))
    _lijst(_alinea(doc, "(geef de aanleiding"), g.get("constateren"))
    _lijst(_alinea(doc, "(geef de strekking"), g.get("veranderen"))

    van_p = _alinea(doc, "Verandert het besluit van")
    toe_p = _alinea(doc, "Voegt het volgende beslispunt")
    weg_p = _alinea(doc, "Verwijderen van het volgende beslispunt")
    of_ps = [x for x in doc.paragraphs if x.text.strip() == "OF"]
    keuze = g.get("wijziging") or "veranderen"
    gekozen = {"veranderen": van_p, "toevoegen": toe_p, "verwijderen": weg_p}[keuze]
    # Alles tussen de gekozen regel en 'Indiener(s)' dat bij de andere keuzes hoort, gaat weg.
    for x in [van_p, toe_p, weg_p] + of_ps:
        if x is not gekozen:
            _weg(x)
    _leeg_tussen(doc, _alinea(doc, "Wijzigt/verandert het besluit"), gekozen, houd=1)
    _leeg_tussen(doc, gekozen, _alinea(doc, "Indiener(s)"), houd=3)
    leeg = copy.deepcopy(gekozen._p)
    _zet(Paragraph(leeg, gekozen._parent), "")
    proto_leeg = Paragraph(leeg, gekozen._parent)

    def citaat(na, regels):
        laatste = na
        for regel in [r for r in regels or [] if str(r).strip()]:
            laatste = _zet(_kloon_na(laatste, proto_leeg), str(regel).strip(), cursief=True, vet=False)
            laatste.paragraph_format.left_indent = Cm(1)
        return laatste

    if keuze == "veranderen":
        _zet(gekozen, "Verandert het besluit van:", vet=True)
        laatste = _kloon_na(gekozen, proto_leeg)
        laatste = citaat(laatste, g.get("van"))
        laatste = _kloon_na(laatste, proto_leeg)
        laatste = _zet(_kloon_na(laatste, proto_leeg), "In:", vet=True, cursief=False)
        laatste = _kloon_na(laatste, proto_leeg)
        citaat(laatste, g.get("in"))
    else:
        _zet(gekozen, "Voegt het volgende beslispunt toe:" if keuze == "toevoegen"
             else "Verwijdert het volgende beslispunt:", vet=True)
        punten = [r for r in g.get("punten") or [] if str(r).strip()]
        if len(punten) > 1:
            _zet(gekozen, "Voegt de volgende beslispunten toe:" if keuze == "toevoegen"
                 else "Verwijdert de volgende beslispunten:", vet=True)
        laatste = _kloon_na(gekozen, proto_leeg)
        citaat(laatste, punten)
    _ruim_lege_lijstpunten(doc)
    _zet_indieners(doc, g)


def _motie(doc, g):
    kop = _alinea(doc, "MOTIE")
    titel = (g.get("titel") or "").strip()
    _vervang(kop, "/MOTIE ", " " if titel else "")
    for r in kop.runs:
        if r.text == "ONDERWERP":
            r.text, r.bold = titel.upper(), True
    behorend = _alinea(doc, "Behorend bij raadsvoorstel")
    vreemd = _alinea(doc, "Nieuw voorstel over een onderwerp")
    of_p = next(x for x in doc.paragraphs if x.text.strip() == "OF")
    if (g.get("vreemd_onderwerp") or "").strip() and not (g.get("raadsvoorstel_titel") or "").strip():
        _vervang(vreemd, "Onderwerp", g["vreemd_onderwerp"].strip())
        _weg(behorend)
    else:
        _vervang(behorend, "titel agendapunt", g.get("raadsvoorstel_titel", ""))
        _vervang(behorend, "nummer", g.get("raadsvoorstel_nummer", ""))
        _weg(vreemd)
    _weg(of_p)
    _vervang(_alinea(doc, "De gemeenteraad van Helmond"), "datum", g.get("vergaderdatum", ""))
    # De opdracht wordt een opsomming in dezelfde stijl als de lijsten erboven.
    lijst_proto = copy.deepcopy(_alinea(doc, "[mening")._p)
    _lijst(_alinea(doc, "[feiten"), g.get("vaststellen"))
    _lijst(_alinea(doc, "[overwegingen"), g.get("overwegingen"))
    _lijst(_alinea(doc, "[mening"), g.get("vinden"))
    opdracht = _alinea(doc, "We vragen het college om")
    zin = {"vragen": "We vragen het college om:", "uitspreken": "We spreken uit dat:",
           "opdracht": "We geven het college opdracht om:"}[g.get("opdracht_soort") or "vragen"]
    _zet(opdracht, zin, vet=True)
    opdracht._p.addnext(lijst_proto)
    _lijst(Paragraph(lijst_proto, opdracht._parent), g.get("dictum"))
    _ruim_lege_lijstpunten(doc)
    _leeg_tussen(doc, _alinea(doc, zin), _alinea(doc, "Indiener(s)"), houd=3)
    _zet_indieners(doc, g)


def _kopvelden(doc, waarden):
    for label, waarde in waarden.items():
        _veld(_alinea(doc, label), waarde)


def _technische_vraag(doc, g):
    namen, partij = _namen(g)
    _kopvelden(doc, {"Onderwerp": g.get("onderwerp", ""), "Ingekomen": "", "Indiener": namen,
                     "Partij": partij, "Portefeu": g.get("portefeuillehouder", ""), "Antwoord": ""})
    start, eind = _alinea(doc, "De vragen"), _alinea(doc, "Antwoorden")
    proto = next(p for p in doc.paragraphs[_index(doc, start) + 1:] if _leeg(p))
    proto = Paragraph(copy.deepcopy(proto._p), proto._parent)
    _leeg_tussen(doc, start, eind, houd=0)
    laatste = _schrijf_vragen(start, proto, g.get("vragen"))
    for _ in range(3):
        laatste = _kloon_na(laatste, proto)


def _vraag_agendapunt(doc, g):
    namen, partij = _namen(g)
    _kopvelden(doc, {"Vergadering": g.get("vergadering", ""), "Agendapunt": g.get("agendapunt", ""),
                     "Ingekomen": "", "Indiener": namen, "Partij": partij,
                     "Portefeuillehouder": g.get("portefeuillehouder", ""), "Datum a": ""})
    start, eind = _alinea(doc, "De vraag"), _alinea(doc, "Het antwoord")
    proto = next(p for p in doc.paragraphs[_index(doc, start) + 1:] if _leeg(p))
    proto = Paragraph(copy.deepcopy(proto._p), proto._parent)
    _leeg_tussen(doc, start, eind, houd=0)
    laatste = _schrijf_vragen(start, proto, g.get("vragen"))
    for _ in range(2):
        laatste = _kloon_na(laatste, proto)


def _politieke_vraag(doc, g):
    namen, partij = _namen(g)
    _kopvelden(doc, {"Onderwerp": g.get("onderwerp", ""), "Indiener": namen, "Partij": partij,
                     "Volgnummer": "", "Ingekomen": "", "Antwoord ": ""})
    aanhef = _alinea(doc, "Geacht college")
    wij = _alinea(doc, "Wij hebben hierover")
    antwoord = _alinea(doc, "Antwoord van het college")
    a, w = _index(doc, aanhef), _index(doc, wij)
    proto_tekst = Paragraph(copy.deepcopy(next(p for p in doc.paragraphs[a + 1:w] if not _leeg(p))._p), aanhef._parent)
    proto_leeg = Paragraph(copy.deepcopy(next(p for p in doc.paragraphs[a + 1:w] if _leeg(p))._p), aanhef._parent)
    # Het voorbeeld van de griffie (EHAD-terrein) gaat eruit, ook de voorbeeldvragen.
    for p in doc.paragraphs[a + 1:w] + doc.paragraphs[w + 1:_index(doc, antwoord)]:
        _weg(p)
    laatste = _kloon_na(aanhef, proto_leeg)
    for alinea in [x for x in g.get("inleiding") or [] if str(x).strip()]:
        laatste = _zet(_kloon_na(laatste, proto_tekst), str(alinea).strip())
        laatste = _kloon_na(laatste, proto_leeg)
    laatste = _schrijf_vragen(wij, proto_leeg, g.get("vragen"), stijl="schuin")
    for _ in range(3):
        laatste = _kloon_na(laatste, proto_leeg)


VULLERS = {"amendement": _amendement, "motie": _motie, "technische_vraag": _technische_vraag,
           "politieke_vraag": _politieke_vraag, "vraag_agendapunt": _vraag_agendapunt}


def standaard(soort):
    with open(os.path.join(MAP, f"{soort}.docx"), "rb") as f:
        return f.read()


def vul(soort, gegevens, sjabloon=None):
    """Geeft het ingevulde Word-bestand (bytes). `sjabloon` is een geüploade versie; anders de standaard."""
    if soort not in VULLERS:
        raise ValueError(f"Onbekende soort: {soort}")
    doc = docx.Document(io.BytesIO(sjabloon or standaard(soort)))
    _plat_velden(doc)
    VULLERS[soort](doc, gegevens or {})
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue()


def controleer(soort, sjabloon):
    """Kan dit (geüploade) sjabloon worden ingevuld? Geeft een foutmelding of None."""
    try:
        vul(soort, VOORBEELD[soort], sjabloon)
        return None
    except Exception as e:  # noqa: BLE001 - de melding gaat naar het scherm
        return str(e) or e.__class__.__name__


def bestandsnaam(soort, gegevens):
    basis = gegevens.get("raadsvoorstel_nummer") or gegevens.get("onderwerp") or gegevens.get("titel") \
        or gegevens.get("vreemd_onderwerp") or gegevens.get("agendapunt") or ""
    basis = re.sub(r"[^\w\- ]+", "", basis).strip().replace(" ", "-")[:60]
    return f"{SOORTEN[soort].split(' (')[0].replace(' ', '-')}{'-' + basis if basis else ''}.docx"
