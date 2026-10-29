"""Fractievergadering (M5): notulen en actielijst uit een transcript (bijvoorbeeld van Notta).

Martijn neemt de vergadering op met Notta en exporteert het transcript (txt, docx, srt of pdf), of plakt de
tekst. Claude maakt daar uitgebreide notulen van: per agendapunt een globale beschrijving van wat is
besproken, de besluiten en de acties, met een besluitenlijst, actielijst en verdeling van de woordvoering.
Twijfelpassages worden apart genoemd. Na vaststellen wordt het transcript verwijderd.
"""
import datetime as dt
import difflib
import html
import io
import json
import re
import threading
import traceback

from . import claude, kennisbank, nieuws, whaller

# De vijf statussen van een actie op het kanbanbord.
STATUSSEN = [("nieuw", "Nieuw"), ("triage", "Triage"), ("behandeling", "In behandeling"),
             ("wachten", "Wachten op"), ("afgerond", "Afgerond")]
STATUS_NAAM = dict(STATUSSEN)

SOORTEN_BESTAND = (".txt", ".md", ".docx", ".srt", ".vtt", ".pdf")
MAX_TRANSCRIPT = 400_000  # tekens; ruim genoeg voor een avond van drie uur

OPDRACHT = """Maak uitgebreide notulen van deze fractievergadering van PRO Helmond.

Werkwijze:
- Volg de agenda als die er is. Bespreekt de fractie iets buiten de agenda, zet het dan bij het passende punt
  (bijvoorbeeld Overige politieke zaken of Rondvraag) of maak een eigen punt.
- Per agendapunt: een globale beschrijving van wat is besproken (meestal één tot drie alinea's: de aanleiding, de
  belangrijkste standpunten en argumenten, en waar de fractie uitkwam), daarna de besluiten en de acties.
  Bij voorbereiding van commissie of raad: per onderwerp wie woordvoerder is en wat de inzet is.
- Een besluit is wat de fractie afspreekt. Een actie heeft een wie en zo mogelijk een wanneer.
  Noem alleen een naam als die uit het transcript blijkt. Weet je het niet, zet dan wie leeg.
- Het transcript komt van spraakherkenning. Namen en woorden kunnen verkeerd zijn gespeld en sprekers heten soms
  "Spreker 1". Gebruik de juiste namen uit de lijst hieronder als duidelijk is wie bedoeld wordt. Is iets
  onduidelijk of tegenstrijdig, zet het dan in 'twijfel' met een korte uitleg. Verzin niets.
- Schrijf zakelijk en neutraal, in de verleden tijd, op taalniveau B2. Geen gedachtestreepjes.
- Persoonlijke opmerkingen of gevoelige uitspraken over personen neem je alleen op als ze nodig zijn voor een
  besluit; zet ze dan zakelijk en kort neer.
- Neem praatjes vooraf, technische problemen met de opname en herhalingen niet op.

Antwoord met alleen één blok ```json met dit object:
{"titel": "Fractievergadering PRO Helmond d-m-jjjj",
 "samenvatting": "drie tot vijf zinnen: de hoofdlijnen van de avond",
 "punten": [{"nummer": "1", "titel": "Mededelingen", "verslag": "alinea's, gescheiden door een lege regel",
             "besluiten": ["..."], "acties": [{"actie": "...", "wie": "naam of leeg", "wanneer": "datum of moment, of leeg"}]}],
 "woordvoering": [{"agendapunt": "...", "vergadering": "Commissie of Raad d-m-jjjj", "woordvoerder": "naam", "standpunt": "inzet in één zin"}],
 "twijfel": ["passage of naam die Martijn moet nagaan, met waarom"]}"""


def tekst_uit_bestand(naam, inhoud):
    """Tekst uit een Notta-export of ander bestand."""
    ext = "." + naam.lower().rsplit(".", 1)[-1] if "." in naam else ""
    if ext in (".srt", ".vtt"):
        return srt_naar_tekst(inhoud.decode("utf-8", errors="replace"))
    if ext in (".txt", ".md", ".docx", ".pdf"):
        return kennisbank._tekst_uit(naam if ext != ".md" else naam, inhoud)  # pylint: disable=protected-access
    raise ValueError(f"Dit bestandstype kan ik niet lezen ({ext or 'onbekend'}). Gebruik txt, docx, srt of pdf.")


def srt_naar_tekst(t):
    """Ondertitels (srt of vtt) naar doorlopende tekst: nummers en tijden weg, opeenvolgende regels van
    dezelfde spreker samen."""
    regels, vorige_spreker = [], None
    for blok in re.split(r"\n\s*\n", t.replace("\r", "")):
        delen = [r.strip() for r in blok.split("\n") if r.strip()]
        delen = [r for r in delen if not re.fullmatch(r"\d+", r) and "-->" not in r and r != "WEBVTT"]
        if not delen:
            continue
        zin = " ".join(delen)
        m = re.match(r"^([^:]{1,40}):\s*(.*)$", zin)
        spreker, zin = (m.group(1), m.group(2)) if m else (None, zin)
        if spreker and spreker == vorige_spreker and regels:
            regels[-1] += " " + zin
        else:
            regels.append(f"{spreker}: {zin}" if spreker else zin)
            vorige_spreker = spreker
    return "\n".join(regels)


def agenda_tekst(datum, inst, invoer, vergaderingen=()):
    """De agenda van de fractievergadering als platte, genummerde tekst (uit het Whaller-bericht)."""
    h = whaller.fractie_bericht(datum, inst, invoer, list(vergaderingen))
    regels, tellers, huidig = [], [], ""
    for stuk in re.split(r"(<[^>]+>)", h):
        tag = stuk.lower()
        if tag.startswith("<ol"):
            if huidig.strip():
                regels.append(huidig.rstrip())
            huidig = ""
            tellers.append(0)
        elif tag == "</ol>":
            if huidig.strip():
                regels.append(huidig.rstrip())
            huidig = ""
            tellers.pop()
        elif tag == "<li>":
            if huidig.strip():
                regels.append(huidig.rstrip())
            tellers[-1] += 1
            huidig = "   " * (len(tellers) - 1) + f"{tellers[-1]}. "
        elif tag in ("</p>", "</h2>"):
            if huidig.strip():
                regels.append(huidig.rstrip())
            huidig = ""
        elif not stuk.startswith("<"):
            huidig += html.unescape(stuk).replace("\xa0", " ")
    if huidig.strip():
        regels.append(huidig.rstrip())
    return "\n".join(r.strip("\n") for r in regels if r.strip())


def _systeem(c):
    profiel = nieuws.instelling(c, "nieuws_profiel", nieuws.STANDAARD_PROFIEL)
    return (f"Je bent de notulist van de fractie van PRO Helmond in de gemeenteraad van Helmond.\n\n{profiel}\n\n"
            f"Vandaag is het {whaller.lang(dt.date.today())}.")


def maak(c, nid):
    n = c.execute("SELECT * FROM notulen WHERE id = %s", (nid,)).fetchone()
    claude.schrijfplafond(c)
    personen = c.execute("SELECT naam, rol FROM personen WHERE actief ORDER BY volgorde, naam").fetchall()
    open_acties = c.execute("""SELECT actie, wie, wanneer FROM fractie_acties WHERE status <> 'afgerond'
                                 AND (notulen_id IS NULL OR notulen_id <> %s) ORDER BY id""", (nid,)).fetchall()
    delen = [f"# Vergadering\nDatum: {whaller.lang(n['datum'])}\n"
             f"Aanwezig: {', '.join(n['aanwezig'] or []) or 'onbekend'}\n"
             f"Afwezig: {', '.join(n['afwezig'] or []) or 'niemand opgegeven'}\n"
             f"Voorzitter: {n['voorzitter'] or 'onbekend'}\nNotulist: {n['notulist'] or 'onbekend'}",
             "# Fractie en steunfractie (juiste spelling van de namen)\n"
             + "\n".join(f"- {p['naam']} ({p['rol']})" for p in personen)]
    if n["agenda"]:
        delen.append("# Agenda\n" + n["agenda"])
    if open_acties:
        delen.append("# Open acties uit eerdere vergaderingen (noem het als de fractie er iets over zegt)\n"
                     + "\n".join(f"- {a['actie']} ({a['wie'] or 'wie onbekend'}{', ' + a['wanneer'] if a['wanneer'] else ''})"
                                 for a in open_acties[:40]))
    delen.append("# Transcript\n" + (n["transcript"] or "")[:MAX_TRANSCRIPT])
    opdracht = OPDRACHT
    if n["aanwijzingen"]:
        opdracht += f"\n\nAanwijzingen van Martijn:\n{n['aanwijzingen']}"
    delen.append("# Opdracht\n" + opdracht)
    model = nieuws.instelling(c, "claude_model_schrijven", "claude-sonnet-5-5")
    antwoord = claude.vraag(c, "notulen", _systeem(c), "\n\n".join(delen), model, max_tokens=16000, timeout=600)
    inhoud = claude.json_uit(antwoord)
    inhoud = json.loads(json.dumps(inhoud, ensure_ascii=False).replace(" — ", ", ").replace("—", ", ").replace(" – ", ", "))
    c.execute("""UPDATE notulen SET status = 'klaar', inhoud = %s, model = %s, fout = NULL, bijgewerkt = now()
                 WHERE id = %s""", (json.dumps(inhoud, ensure_ascii=False), model, nid))
    # De actielijst opnieuw opbouwen voor deze notulen. Wat Martijn afvinkte of zelf aanpaste blijft staan;
    # een nieuwe actie die daar sterk op lijkt, komt er niet nog een keer bij.
    c.execute("DELETE FROM fractie_acties WHERE notulen_id = %s AND status = 'nieuw' AND NOT bewerkt", (nid,))
    blijft = [_actie_sleutel(r["actie"]) for r in
              c.execute("SELECT actie FROM fractie_acties WHERE notulen_id = %s", (nid,)).fetchall()]
    for p in inhoud.get("punten", []):
        for a in p.get("acties") or []:
            k = _actie_sleutel(a.get("actie") or "")
            if k and difflib.get_close_matches(k, blijft, n=1, cutoff=0.8):
                continue
            if (a.get("actie") or "").strip():
                c.execute("""INSERT INTO fractie_acties (notulen_id, actie, wie, wanneer, punt)
                             VALUES (%s, %s, %s, %s, %s)""",
                          (nid, a["actie"].strip(), (a.get("wie") or "").strip(), (a.get("wanneer") or "").strip(),
                           p.get("titel") or ""))


def _actie_sleutel(t):
    return re.sub(r"[^a-z0-9]+", " ", (t or "").lower()).strip()


def start(pool, nid):
    def werk():
        with pool.connection() as c:
            try:
                maak(c, nid)
            except Exception as e:  # noqa: BLE001
                c.execute("UPDATE notulen SET status = 'fout', fout = %s, bijgewerkt = now() WHERE id = %s",
                          (f"{e}\n\n{traceback.format_exc()[-1200:]}", nid))
    threading.Thread(target=werk, daemon=True).start()


# ---------- Word ----------

def word(n, acties):
    """Notulen als Word-bestand, volgens het vaste format (bijlage B), uitgebreid."""
    import docx  # pylint: disable=import-outside-toplevel
    from docx.enum.text import WD_COLOR_INDEX  # pylint: disable=import-outside-toplevel
    from docx.shared import Pt  # pylint: disable=import-outside-toplevel
    i = n["inhoud"] or {}
    doc = docx.Document()
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(11)
    doc.add_heading(i.get("titel") or f"Fractievergadering PRO Helmond {whaller.d(n['datum'])}", level=1)
    if n["status"] != "vastgesteld":
        p = doc.add_paragraph()
        r = p.add_run("Concept, nog niet vastgesteld. Intern: alleen voor de fractie.")
        r.italic = True
    kop = doc.add_table(rows=0, cols=2)
    kop.style = "Table Grid"
    for label, waarde in (("Datum", whaller.lang(n["datum"])), ("Plaats", n["plaats"] or ""),
                          ("Aanwezig", ", ".join(n["aanwezig"] or [])), ("Afwezig", ", ".join(n["afwezig"] or [])),
                          ("Voorzitter", n["voorzitter"] or ""), ("Notulist", n["notulist"] or "")):
        cellen = kop.add_row().cells
        cellen[0].text, cellen[1].text = label, waarde
    if i.get("samenvatting"):
        doc.add_heading("Samenvatting", level=2)
        doc.add_paragraph(i["samenvatting"])
    for p in i.get("punten", []):
        doc.add_heading(f"{p.get('nummer', '')}. {p.get('titel', '')}".strip(". "), level=2)
        for alinea in re.split(r"\n\s*\n", p.get("verslag") or ""):
            if alinea.strip():
                doc.add_paragraph(alinea.strip())
        if p.get("besluiten"):
            doc.add_paragraph().add_run("Besluit" + ("en" if len(p["besluiten"]) > 1 else "")).bold = True
            for b in p["besluiten"]:
                doc.add_paragraph(b, style="List Bullet")
        if p.get("acties"):
            doc.add_paragraph().add_run("Acties").bold = True
            for a in p["acties"]:
                doc.add_paragraph(_actie_regel(a), style="List Bullet")
    besluiten = [b for p in i.get("punten", []) for b in (p.get("besluiten") or [])]
    if besluiten:
        doc.add_heading("Besluitenlijst", level=2)
        for b in besluiten:
            doc.add_paragraph(b, style="List Number")
    doc.add_heading("Actielijst", level=2)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for cel, kop_tekst in zip(t.rows[0].cells, ("Nr", "Actie", "Wie", "Wanneer", "Status")):
        cel.text = kop_tekst
    for nr, a in enumerate(acties, 1):
        cellen = t.add_row().cells
        for cel, waarde in zip(cellen, (str(nr), a["actie"], a["wie"] or "", a["wanneer"] or "",
                                        STATUS_NAAM.get(a["status"], a["status"]))):
            cel.text = waarde
    if i.get("woordvoering"):
        doc.add_heading("Verdeling woordvoering", level=2)
        w = doc.add_table(rows=1, cols=4)
        w.style = "Table Grid"
        for cel, kop_tekst in zip(w.rows[0].cells, ("Agendapunt", "Vergadering", "Woordvoerder", "Standpunt")):
            cel.text = kop_tekst
        for r in i["woordvoering"]:
            cellen = w.add_row().cells
            for cel, k in zip(cellen, ("agendapunt", "vergadering", "woordvoerder", "standpunt")):
                cel.text = r.get(k) or ""
    if i.get("twijfel") and n["status"] != "vastgesteld":
        doc.add_heading("Na te gaan voor vaststelling", level=2)
        for tw in i["twijfel"]:
            run = doc.add_paragraph(style="List Bullet").add_run(tw)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue()


def _actie_regel(a):
    extra = ", ".join(x for x in (a.get("wie"), a.get("wanneer")) if x)
    return a.get("actie", "") + (f" ({extra})" if extra else "")


def bestandsnaam(n):
    return f"{n['datum']:%Y-%m-%d}-notulen-fractievergadering{'' if n['status'] == 'vastgesteld' else '-concept'}.docx"
