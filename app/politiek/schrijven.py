"""Bijdragenatelier (M4, M8): concepten schrijven met Claude.

Een concept begint bij een agendapunt uit Notubiz, een stuk uit de Inbox, een eigen tekst of een idee.
Promonitor zoekt de stukken erbij (raadsvoorstel en bijlagen, gelezen als tekst), haalt passende passages
uit de kennisbank (verkiezingsprogramma, coalitieakkoord, begroting, eerdere bijdragen) en laat Claude
schrijven in de huisstijl van PRO Helmond. Moties, amendementen en vragen komen in het format van de griffie.
Elk concept heeft een lijst met feiten en hun bron, zodat Martijn kan nagaan wat klopt.
"""
import datetime as dt
import io
import json
import re
import threading
import traceback

from . import agenda, claude, kennisbank, nieuws, notubiz, raadsinfo, sjablonen, whaller

SOORTEN = {
    "analyse": "Analyse van een agendapunt",
    "technische_vraag": "Technische vragen",
    "vraag_agendapunt": "Vragen bij een geagendeerd onderwerp",
    "bijdrage": "Bijdrage (spreektekst) voor commissie of raad",
    "amendement": "Amendement",
    "motie": "Motie",
    "politieke_vraag": "Politieke vragen (schriftelijke vragen aan het college)",
    "stemverklaring": "Stemverklaring",
    "feedback": "Feedback op een tekst van een fractiegenoot",
    "analyse_partij": "Analyse van een bijdrage van een andere partij",
}
# Soorten met een eigen uitvoer (geen Value-Villain-Vision, wel opmerkingen of punten).
EIGEN_UITVOER = {"feedback", "analyse_partij"}
WOORDEN_PER_MINUUT = 130
MET_SJABLOON = {"technische_vraag", "vraag_agendapunt", "amendement", "motie", "politieke_vraag"}
VOORKEUR = ("verkiezingsprogramma-groenlinks-pvda", "coalitieakkoord-2026", "programmabegroting-2027", "samenvatting-")

HUISSTIJL = """Huisstijl van PRO Helmond (volg die altijd):
- Taalniveau B2: korte zinnen, alledaagse woorden, actieve vorm. Leg vakjargon uit bij het eerste gebruik.
- Gebruik nooit gedachtestreepjes (het lange streepje). Gebruik een dubbele punt, een komma, haakjes of een nieuwe zin.
- Toon: constructief-kritisch. Steun het doel waar dat kan, richt de kritiek op de uitvoering. Niet populistisch.
- Value, Villain, Vision: begin bij een gedeelde waarde die iedereen in de zaal kan onderschrijven (niet bij het
  probleem), benoem dan concreet wat die waarde in de weg staat (een regel, keuze of werkwijze, nooit een persoon) en
  sluit af met een concreet, hoopvol beeld van hoe het eruitziet als het goed geregeld is, gekoppeld aan jouw vragen
  en verzoeken. Bij een puur technisch voorstel mag dit licht blijven.
- Zoek de spanning: het belang van de inwoner tegenover het organisatiebelang, ontbrekende cijfers, de kanttekeningen
  die het college zelf noemt, grip van de raad, overgangsrecht, en of de titel de lading dekt.
- Feiten: alleen wat in de aangeleverde stukken of de kennisbank staat. Verzin geen aantallen, data of citaten.
  Twijfel je, zeg dat dan en zet het in de feitenlijst als 'controleren'.
- FvD wordt nergens genormaliseerd of als bondgenoot neergezet."""

OPDRACHT = {
    "analyse": """Maak een analyse van dit agendapunt voor de fractie. Opbouw in Markdown:
## In het kort (drie tot vijf zinnen: wat stelt het college voor en wat moet de raad besluiten)
## Wat betekent dit voor Helmond
## Toets aan ons programma, het coalitieakkoord en de begroting
## Zwakke plekken en ontbrekende informatie
## Kansen voor PRO Helmond
## Advies: standpunt, vragen en eventueel een motie of amendement
Wees concreet en verwijs naar de stukken (paginanummer als dat kan).""",
    "bijdrage": """Schrijf een spreektekst van twee tot drie minuten. Aanhef "Voorzitter," en afsluiting "Dank u wel."
Opbouw: waarde (waar dit echt over gaat), wat wij steunen, de belemmering (waar het misgaat), onze vragen (genummerd en
aangekondigd), concrete verzoeken of toezeggingen, het perspectief (hoe het er dan uitziet) en een voorwaardelijke
conclusie ("krijgen wij die toezeggingen, dan ...; zo niet, dan overwegen wij een amendement").
Begin het document met drie regels: PRO Helmond, de vergadering met datum, en het raadsvoorstel met nummer en onderwerp.""",
    "stemverklaring": """Schrijf een korte stemverklaring (hooguit één minuut): hoe stemt PRO Helmond en waarom, met de waarde
voorop, de kern van onze zorg of steun en wat we van het college verwachten. Aanhef "Voorzitter," en afsluiting "Dank u wel.\"""",
    "feedback": """Geef feedback op de tekst van een fractiegenoot (een bijdrage, vragen, motie of stuk). Eerlijk en
vriendelijk: het is een collega. Toets aan de huisstijl (B2, geen gedachtestreepjes, Value-Villain-Vision), aan het
programma en aan de stukken.
Schrijf eerst een korte feedbacknotitie in Markdown:
## Kort oordeel (twee of drie zinnen)
## Drie verbeterpunten (concreet, met een voorbeeldzin)
## Feiten om te controleren
## Risico's (wat kan tegen ons worden gebruikt, wat kan de wethouder makkelijk pareren)
## Leesbaarheid en spreektijd
## Tekstvoorstellen (alleen de zinnen of alinea's die je zou herschrijven)""",
    "analyse_partij": """Analyseer de bijdrage van een andere partij (uit een verslag, een videoregistratie, aantekeningen of hun
eigen tekst). Opbouw in Markdown:
## Wat zegt de partij (kort en feitelijk)
## Afwijkingen (van hun eigen programma of eerdere standpunten als die bekend zijn, en van het coalitieakkoord als het
een coalitiepartij is; met bron, en zeg het als je hun programma niet hebt)
## Zwakke plekken en kansen voor een interruptie of vraag
## Kansen voor samenwerking met PRO Helmond (nooit FvD als bondgenoot)
## Punten voor de bewakingsmonitor (toezeggingen, beloften, uitspraken om op terug te komen)""",
    "technische_vraag": "Stel technische vragen aan de ambtelijke organisatie: feitelijke vragen over het voorstel en de stukken, geen politieke. Groepeer ze met tussenkoppen (## kop) als het er meer dan zes zijn. Vraag naar ontbrekende cijfers, onderbouwing, risico's en alternatieven.",
    "vraag_agendapunt": "Stel vragen bij dit geagendeerde onderwerp aan het college, scherp en kort, die de bespreking voorbereiden.",
    "politieke_vraag": "Schrijf schriftelijke vragen aan het college (artikel 40 van het Reglement van Orde): een korte inleiding (wat is er aan de hand, met bron) en daarna genummerde vragen. Politiek scherp maar feitelijk.",
    "amendement": "Schrijf een amendement: een voorstel om het ontwerpbesluit te wijzigen, zo geformuleerd dat het direct in het besluit kan. Constateringen feitelijk met bron, de wijziging precies (welke tekst verandert in welke).",
    "motie": "Schrijf een motie: constateringen, overwegingen, wat we vinden en een concrete opdracht of uitspraak. Als er geen raadsvoorstel bij hoort, is het een motie vreemd aan de orde van de dag.",
}


def _velden_uitleg(soort):
    regels = []
    for naam, label, aard in sjablonen.VELDEN[soort]:
        if aard == "regels":
            vorm = "lijst van korte teksten"
        elif aard == "tekst":
            vorm = "tekst"
        else:
            vorm = "een van: " + ", ".join(k for k, _ in aard[1])
        regels.append(f'  "{naam}": {vorm}  ({label})')
    return "{\n" + ",\n".join(regels) + "\n}"


def systeem(c):
    profiel = nieuws.instelling(c, "nieuws_profiel", nieuws.STANDAARD_PROFIEL)
    return (f"Je bent de schrijver en analist van Promonitor, de politieke cockpit van een raadslid in Helmond.\n\n"
            f"{profiel}\n\n{HUISSTIJL}\n\nVandaag is het {whaller.lang(dt.date.today())}.")


def uitvoerformaat(soort):
    if soort == "feedback":
        return ("Zet na de notitie één blok ```json met:\n"
                '{"titel": "Feedback op ... (korte titel)", "oordeel": "één zin",\n'
                ' "opmerkingen": [{"citaat": "letterlijke tekst uit de tekst van de collega, 4 tot 15 woorden, precies zoals hij er staat",\n'
                '                  "opmerking": "wat beter kan en waarom", "voorstel": "nieuwe formulering, of leeg"}],\n'
                ' "feiten": [{"bewering": "...", "bron": "document en pagina", "status": "bron" of "controleren"}],\n'
                ' "toelichting": "twee zinnen aan Martijn"}\n'
                "Geef zes tot vijftien opmerkingen, verspreid over de hele tekst. Deze komen in de kantlijn van het Word-bestand.")
    if soort == "analyse_partij":
        return ("Zet na de tekst één blok ```json met:\n"
                '{"titel": "Analyse bijdrage ... (partij en onderwerp)", "partij": "naam",\n'
                ' "bewaken": ["uitspraak of belofte om op terug te komen, met datum en vergadering als bekend"],\n'
                ' "feiten": [{"bewering": "...", "bron": "document en pagina", "status": "bron" of "controleren"}],\n'
                ' "toelichting": "twee zinnen aan Martijn"}')
    if soort in MET_SJABLOON:
        return ("Antwoord met alleen één blok ```json met dit object:\n"
                '{"titel": "korte titel", "velden": ' + _velden_uitleg(soort) + ',\n'
                ' "feiten": [{"bewering": "...", "bron": "document en pagina", "status": "bron" of "controleren"}],\n'
                ' "slordigheden": ["typefouten, tegenstrijdige data of ontbrekende bijlagen in de stukken"],\n'
                ' "toelichting": "twee of drie zinnen aan Martijn: wat je koos en waarom"}\n'
                "Indieners: Martijn Buurman, tenzij anders gevraagd. Gebruik in de vragen '## kop' voor tussenkoppen en '- ' voor deelvragen.")
    return ("Schrijf eerst de tekst in Markdown. Zet daarna één blok ```json met:\n"
            '{"titel": "korte titel", "vvv": {"value": "één zin", "villain": "één zin", "vision": "één zin"},\n'
            ' "feiten": [{"bewering": "...", "bron": "document en pagina", "status": "bron" of "controleren"}],\n'
            ' "slordigheden": ["typefouten, tegenstrijdige data of ontbrekende bijlagen in de stukken"],\n'
            ' "toelichting": "twee of drie zinnen aan Martijn: wat je koos en waarom"}')


# ---------- stukken en context ----------

def tekst_van_document(c, doc):
    """Tekst van een Notubiz-document uit de kennisbank, of ophalen en bewaren."""
    pad = f"notubiz/{doc['id']}"
    r = c.execute("SELECT id, tekst FROM documenten WHERE pad = %s", (pad,)).fetchone()
    if r:
        return r["id"], r["tekst"]
    tekst = raadsinfo.pdf_tekst(notubiz.document_ophalen(doc["url"])).replace("\x00", "")
    r = c.execute("""INSERT INTO documenten (titel, soort, pad, tekst, woorden, bron_bestand)
                     VALUES (%s, 'raadsstuk', %s, %s, %s, %s)
                     ON CONFLICT (pad) DO UPDATE SET tekst = EXCLUDED.tekst RETURNING id""",
                  (doc["titel"], pad, tekst, len(tekst.split()), doc["url"])).fetchone()
    return r["id"], tekst


def _volgorde(doc):
    t = doc["titel"].lower()
    if re.match(r"^\s*(rv|raadsvoorstel)\b", t):
        return 0
    if "besluit" in t:
        return 1
    return 2


def stukken_tekst(c, documenten, maximum=60000):
    """De stukken bij een agendapunt als tekst: raadsvoorstel eerst en ruim, bijlagen ingekort."""
    delen, ids, totaal, gemist = [], [], 0, []
    for doc in sorted(documenten, key=_volgorde):
        try:
            did, tekst = tekst_van_document(c, doc)
        except Exception as e:  # noqa: BLE001
            gemist.append(f"{doc['titel']} (niet te downloaden: {str(e)[:80]})")
            continue
        ids.append(did)
        ruimte = 30000 if _volgorde(doc) == 0 else 8000
        ruimte = min(ruimte, maximum - totaal)
        if ruimte < 1000:
            gemist.append(f"{doc['titel']} (geen ruimte meer)")
            continue
        stuk = tekst[:ruimte] + ("\n[... ingekort]" if len(tekst) > ruimte else "")
        delen.append(f"## Document: {doc['titel']}\n{stuk}")
        totaal += len(stuk)
    return "\n\n".join(delen), ids, gemist


def agendapunten(van=None, tot=None):
    """Agendapunten van commissie en raad in de komende weken, met hun documenten (voor de keuzelijst)."""
    van = van or dt.date.today() - dt.timedelta(days=3)
    tot = tot or dt.date.today() + dt.timedelta(days=42)
    uit = []
    for ev in notubiz.evenementen(van, tot):
        if ev.get("alleen_leden") or ev["soort"] not in ("commissie", "raad"):
            continue
        try:
            delen = notubiz.commissieavond(ev["id"])["commissies"] if ev["type"] == "assembly" else [notubiz.vergadering(ev["id"])]
        except Exception:  # noqa: BLE001
            continue
        for verg in delen:
            for p in verg["agenda"]:
                if p["standaard"] or p["diepte"] > 0 or agenda.is_formeel(p["onderwerp"]):
                    continue
                uit.append({"vergadering": f"{verg['naam']} {whaller.d(verg['start'].date())}",
                            "datum": verg["start"].date().isoformat(), "onderwerp": p["onderwerp"],
                            "portefeuillehouder": p["portefeuillehouder"], "nummer": p["nummer"],
                            "sectie": p["sectie"], "punt_id": p["id"],
                            "documenten": [{"id": d["id"], "titel": d["titel"], "url": d["url"]} for d in p["documenten"]]})
    return uit


# ---------- schrijven ----------

def _splits(antwoord):
    """Markdown-tekst en het JSON-blok uit een antwoord halen."""
    m = list(re.finditer(r"```json\s*(.*?)```", antwoord, re.S))
    if m:
        laatste = m[-1]
        tekst = antwoord[:laatste.start()].strip()
        try:
            return tekst, json.loads(laatste.group(1))
        except json.JSONDecodeError:
            pass
    try:
        return "", claude.json_uit(antwoord)
    except Exception:  # noqa: BLE001
        return antwoord.strip(), {}


def _geen_gedachtestreepjes(t):
    return (t or "").replace(" — ", ", ").replace("—", ", ").replace(" – ", ", ")


def maak_concept(c, concept_id):
    """Schrijft een concept (draait in de achtergrond). Werkt de rij in concepten bij."""
    k = c.execute("SELECT * FROM concepten WHERE id = %s", (concept_id,)).fetchone()
    bron, invoer, soort = k["bron"] or {}, k["invoer"] or {}, k["soort"]
    claude.schrijfplafond(c)

    delen, gebruikt, gemist = [], [], []
    if bron.get("type") == "agendapunt":
        delen.append(f"# Agendapunt\nVergadering: {bron.get('vergadering')}\nOnderwerp: {bron.get('onderwerp')}\n"
                     f"Behandeling: {bron.get('sectie') or '-'}\nPortefeuillehouder: {bron.get('portefeuillehouder') or '-'}")
        tekst, gebruikt, gemist = stukken_tekst(c, bron.get("documenten", []))
        if tekst:
            delen.append("# De stukken\n" + tekst)
    elif bron.get("type") == "stuk":
        s = c.execute("SELECT * FROM stukken WHERE id = %s", (bron["stuk_id"],)).fetchone()
        if s:
            delen.append(f"# Aanleiding: {raadsinfo.SOORTEN.get(s['soort'], s['soort'])}\n{s['titel']}\n"
                         f"{s['bronnaam'] or ''} {s['datum'] or ''}\n{s['context'] or ''}\n{s['reden'] or ''}\nLink: {s['document_url']}")
            if s["document_id"] and s["bron"] == "notubiz":
                tekst, gebruikt, gemist = stukken_tekst(c, [{"id": s["document_id"], "titel": s["titel"], "url": s["document_url"]}])
                delen.append("# Het stuk\n" + tekst)
    if invoer.get("tekst"):
        kop = "Tekst om mee te werken"
        if soort == "feedback":
            st = _spreektijd(soort, invoer)
            kop = f"De tekst van de collega ({st['woorden']} woorden, ongeveer {st['spreektijd']} minuut spreektijd)"
        elif soort == "analyse_partij":
            kop = "De bijdrage van de andere partij"
        delen.append(f"# {kop}\n" + invoer["tekst"][:30000])
    if invoer.get("vorige"):
        delen.append("# Vorige versie van dit concept\n" + invoer["vorige"][:30000])
    zoek = " ".join(filter(None, [bron.get("onderwerp"), k["titel"], invoer.get("aanwijzingen", "")[:400]]))
    kb = kennisbank.context(lambda sql, p=None: c.execute(sql, p).fetchall(), zoek, uitsluiten=gebruikt, voorkeur=VOORKEUR)
    if kb:
        delen.append("# Uit de kennisbank (programma, coalitieakkoord, begroting, eerdere bijdragen en raadsstukken)\n" + kb)
    opdracht = f"# Opdracht: {SOORTEN[soort]}\n{OPDRACHT[soort]}"
    if invoer.get("aanwijzingen"):
        opdracht += f"\n\nAanwijzingen van Martijn:\n{invoer['aanwijzingen']}"
    if invoer.get("vorige"):
        opdracht += "\n\nHerschrijf de vorige versie volgens de aanwijzingen. Houd wat goed was."
    delen.append(opdracht + "\n\n" + uitvoerformaat(soort))

    model = nieuws.instelling(c, "claude_model_schrijven", "claude-sonnet-5-5")
    antwoord = claude.vraag(c, "schrijven", systeem(c), "\n\n".join(delen), model, max_tokens=8000, timeout=400)
    tekst, meta = _splits(antwoord)
    tekst = _geen_gedachtestreepjes(tekst)
    velden = meta.get("velden") if soort in MET_SJABLOON else None
    if velden:
        velden = json.loads(_geen_gedachtestreepjes(json.dumps(velden, ensure_ascii=False)))
        tekst = tekst or velden_als_tekst(soort, velden)
    c.execute("""UPDATE concepten SET status = 'klaar', tekst = %s, velden = %s, meta = %s, model = %s,
                   titel = coalesce(nullif(%s, ''), titel), bijgewerkt = now(), fout = NULL WHERE id = %s""",
              (tekst, json.dumps(velden) if velden else None,
               json.dumps({**meta, "gemist": gemist, "documenten": len(gebruikt), **_spreektijd(soort, invoer)},
                          ensure_ascii=False), model,
               (meta.get("titel") or "")[:200], concept_id))


def _vraagregel(w):
    w = str(w)
    if w.startswith("## "):
        return f"\n### {w[3:]}"
    if w.startswith("- "):
        return f"    - {w[2:]}"
    return f"1. {w}"


def _spreektijd(soort, invoer):
    """Bij feedback: lengte en geschatte spreektijd van de tekst van de collega (uitgerekend, niet geschat door Claude)."""
    if soort != "feedback" or not invoer.get("tekst"):
        return {}
    woorden = len(re.findall(r"\w+", invoer["tekst"]))
    return {"woorden": woorden, "spreektijd": round(woorden / WOORDEN_PER_MINUUT, 1)}


def velden_als_tekst(soort, velden):
    """Leesbare weergave van een ingevuld sjabloon (voor het scherm)."""
    regels = []
    for naam, label, aard in sjablonen.VELDEN[soort]:
        waarde = velden.get(naam)
        if not waarde:
            continue
        label = label.split(" (")[0]
        if isinstance(waarde, list):
            if naam == "vragen":
                regels.append(f"**{label}**\n" + "\n".join(_vraagregel(w) for w in waarde))
            else:
                regels.append(f"**{label}**\n" + "\n".join(f"- {w}" for w in waarde))
        else:
            if aard not in ("tekst", "regels"):
                waarde = dict(aard[1]).get(waarde, waarde)
            regels.append(f"**{label}:** {waarde}")
    return "\n\n".join(regels)


def start(pool, concept_id):
    """Start het schrijven in de achtergrond."""
    def werk():
        with pool.connection() as c:
            try:
                maak_concept(c, concept_id)
            except Exception as e:  # noqa: BLE001
                c.execute("UPDATE concepten SET status = 'fout', fout = %s, bijgewerkt = now() WHERE id = %s",
                          (f"{e}\n\n{traceback.format_exc()[-1200:]}", concept_id))
    threading.Thread(target=werk, daemon=True).start()


# ---------- Word ----------

def _runs(alinea, tekst):
    for i, deel in enumerate(re.split(r"(\*\*[^*]+\*\*)", tekst)):
        if deel.startswith("**") and deel.endswith("**"):
            alinea.add_run(deel[2:-2]).bold = True
        elif deel:
            alinea.add_run(deel)


def word_uit_markdown(markdown, titel=""):
    import docx  # pylint: disable=import-outside-toplevel
    from docx.shared import Pt  # pylint: disable=import-outside-toplevel
    doc = docx.Document()
    stijl = doc.styles["Normal"]
    stijl.font.name, stijl.font.size = "Calibri", Pt(11)
    for regel in (markdown or "").replace("\r", "").split("\n"):
        r = regel.rstrip()
        if not r.strip():
            continue
        if r.startswith("### "):
            doc.add_heading(r[4:], level=3)
        elif r.startswith("## "):
            doc.add_heading(r[3:], level=2)
        elif r.startswith("# "):
            doc.add_heading(r[2:], level=1)
        elif re.match(r"^\s*[-*] ", r):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", r))
        elif re.match(r"^\s*\d+[.)] ", r):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)] ", "", r))
        else:
            _runs(doc.add_paragraph(), r)
    if titel:
        doc.core_properties.title = titel
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue()


def word(c, k):
    """Het Word-bestand van een concept: in het griffie-format als dat bestaat, anders uit de tekst."""
    if k["soort"] in MET_SJABLOON and k["velden"]:
        eigen = c.execute("SELECT inhoud FROM sjablonen WHERE soort = %s", (k["soort"],)).fetchone()
        return sjablonen.vul(k["soort"], k["velden"], bytes(eigen["inhoud"]) if eigen else None)
    return word_uit_markdown(k["tekst"], k["titel"])


def _norm(t):
    return re.sub(r"\s+", " ", (t or "").replace("\u2019", "'").replace("\u2018", "'")).strip().lower()


def word_met_opmerkingen(k, origineel=None):
    """De tekst van de collega met de opmerkingen van Claude in de kantlijn (Word-opmerkingen).

    origineel: het geüploade Word-bestand; dan blijft de opmaak van de collega staan. Anders maken we een
    Word-bestand van de geplakte tekst. Geeft (bytes, aantal geplaatste opmerkingen) terug."""
    import difflib  # pylint: disable=import-outside-toplevel
    import docx  # pylint: disable=import-outside-toplevel
    if origineel:
        doc = docx.Document(io.BytesIO(origineel))
    else:
        doc = docx.Document()
        for alinea in re.split(r"\n\s*\n|\n", (k["invoer"] or {}).get("tekst", "")):
            if alinea.strip():
                doc.add_paragraph(alinea.strip())
    alineas = [p for p in doc.paragraphs if p.text.strip() and p.runs]
    for t in doc.tables:
        for rij in t.rows:
            for cel in rij.cells:
                alineas += [p for p in cel.paragraphs if p.text.strip() and p.runs]
    geplaatst = 0
    for o in (k["meta"] or {}).get("opmerkingen") or []:
        citaat = _norm(o.get("citaat"))
        if not citaat:
            continue
        doel = next((p for p in alineas if citaat in _norm(p.text)), None)
        if doel is None:
            # Geen letterlijke match (Claude citeerde net anders): de alinea die er het meest op lijkt.
            scores = [(difflib.SequenceMatcher(None, citaat, _norm(p.text)).find_longest_match(
                0, len(citaat), 0, len(_norm(p.text))).size, p) for p in alineas]
            beste = max(scores, key=lambda x: x[0], default=(0, None))
            doel = beste[1] if beste[0] >= max(12, len(citaat) // 2) else None
        tekst = o.get("opmerking") or ""
        if o.get("voorstel"):
            tekst += f"\nVoorstel: {o['voorstel']}"
        if doel is None:
            doel = alineas[0] if alineas else None
            tekst = f"(Bij: \"{o.get('citaat')}\") " + tekst
        if doel is not None and tekst.strip():
            doc.add_comment(doel.runs, text=_geen_gedachtestreepjes(tekst.strip()), author="Promonitor", initials="PM")
            geplaatst += 1
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue(), geplaatst


def bestandsnaam(k):
    basis = re.sub(r"[^\w\- ]+", "", k["titel"] or SOORTEN[k["soort"]]).strip().replace(" ", "-")[:70]
    return f"{dt.date.today():%Y-%m-%d}-{basis or k['soort']}.docx"


def md_html(tekst):
    """Eenvoudige Markdown naar HTML voor het scherm (koppen, lijsten met ingesprongen deelpunten, vet, alinea's)."""
    import html as _html  # pylint: disable=import-outside-toplevel
    uit, lijst, sub, nummer = [], None, False, 0

    def sluit_sub():
        nonlocal sub
        if sub:
            uit.append("</ul></li>")
            sub = False

    for regel in (tekst or "").replace("\r", "").split("\n"):
        r = regel.rstrip()
        m_ul, m_ol = re.match(r"^(\s*)[-*] (.*)", r), re.match(r"^\s*\d+[.)] (.*)", r)
        soort = "ul" if m_ul else "ol" if m_ol else None
        inhoud = _html.escape(m_ul.group(2) if m_ul else m_ol.group(1) if m_ol else r.lstrip("# ").strip())
        inhoud = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", inhoud)
        if m_ul and m_ul.group(1) and lijst:
            # Ingesprongen deelpunt onder het vorige punt.
            if not sub:
                uit[-1] = uit[-1][:-5]  # </li> weghalen
                uit.append("<ul>")
                sub = True
            uit.append(f"<li>{inhoud}</li>")
            continue
        if not r.strip():
            continue
        sluit_sub()
        if lijst and soort != lijst:
            uit.append(f"</{lijst}>")
            lijst = None
        if soort:
            if not lijst:
                # Een genummerde lijst die na een kop doorloopt, telt door (zoals in het Word-format).
                uit.append(f'<ol start="{nummer + 1}">' if soort == "ol" else "<ul>")
                lijst = soort
            if soort == "ol":
                nummer += 1
            uit.append(f"<li>{inhoud}</li>")
        elif r.startswith("#"):
            niveau = min(4, len(r) - len(r.lstrip("#")) + 1)
            uit.append(f"<h{niveau}>{inhoud}</h{niveau}>")
        else:
            uit.append(f"<p>{inhoud}</p>")
    sluit_sub()
    if lijst:
        uit.append(f"</{lijst}>")
    return "\n".join(uit)
