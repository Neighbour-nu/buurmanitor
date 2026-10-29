"""Begrotingshulp: bij voorjaarsnota, begroting en jaarrekening de verschillen met eerdere jaren en met beloften.

Claude zoekt in de gekozen documenten de bedragen bij een thema en geeft ze met bron (document en pagina) en
precies zoals ze in het stuk staan. Promonitor rekent zelf de verschillen uit: tussen de jaren in één document
(meerjarig) en tussen documenten voor hetzelfde jaar. Daarnaast toetst Claude de beloften uit het coalitieakkoord
en ons programma aan wat er in de stukken staat.
"""
import datetime as dt
import io
import json
import re
import threading
import traceback

from . import claude, kennisbank, nieuws, schrijven, whaller

DREMPEL_PROCENT = 10       # een verschil vanaf zoveel procent is opvallend
DREMPEL_EURO = 100_000     # of vanaf zoveel euro

OPDRACHT = """Zoek in de passages hierboven de bedragen die horen bij het thema, voor een raadslid van de oppositie.
Regels:
- Neem alleen bedragen over die letterlijk in de passages staan. Verzin niets en reken niet zelf: dat doet Promonitor.
- Geef per bedrag het document, de pagina (als die erbij staat) en de tekst precies zoals in het stuk ("bedrag_bron").
- Zet het bedrag om naar hele euro's in "bedrag": staat er in de tabel "x € 1.000", vermenigvuldig dan met 1000 en zet
  dat in "eenheid_bron". Lasten en uitgaven positief, baten en inkomsten ook positief maar met "soort": "baten".
- Gebruik dezelfde naam voor dezelfde post in verschillende jaren en documenten, anders kan Promonitor niet vergelijken.
- Beloften: wat beloven het coalitieakkoord en ons verkiezingsprogramma over dit thema, en staat er geld voor in de stukken?
- Schrijf op taalniveau B2, zonder gedachtestreepjes.

Antwoord met alleen één blok ```json:
{"titel": "korte titel",
 "posten": [{"post": "naam van de post", "programma": "programma of taakveld",
             "bedragen": [{"jaar": "2027", "bedrag": 125000, "bedrag_bron": "125", "eenheid_bron": "x € 1.000",
                           "soort": "lasten|baten|investering|reserve|saldo", "document": "titel", "pagina": "p. 12"}]}],
 "beloften": [{"belofte": "...", "bron": "document en pagina", "in_stukken": "ja|deels|nee|onduidelijk", "toelichting": "..."}],
 "opvallend": ["wat valt op, met bron"],
 "vragen": ["technische vraag die de raad zou moeten stellen"],
 "toelichting": "twee of drie zinnen aan Martijn: wat je vond en wat je niet kon vinden"}"""


def _bedrag(x):
    try:
        return float(str(x).replace("€", "").replace(" ", "").replace(".", "").replace(",", ".")) if isinstance(x, str) else float(x)
    except (TypeError, ValueError):
        return None


def bereken(posten):
    """Verschillen per post: tussen opeenvolgende jaren in hetzelfde document, en tussen documenten in hetzelfde jaar."""
    rijen = []
    for p in posten:
        bedragen = [b for b in p.get("bedragen") or [] if _bedrag(b.get("bedrag")) is not None]
        for b in bedragen:
            b["bedrag"] = _bedrag(b["bedrag"])
        sleutel = lambda b: (b.get("soort") or "", b.get("document") or "")  # noqa: E731
        groepen = {}
        for b in bedragen:
            groepen.setdefault(sleutel(b), []).append(b)
        for (soort, doc), groep in groepen.items():
            groep.sort(key=lambda b: str(b.get("jaar")))
            for a, b in zip(groep, groep[1:]):
                rijen.append(_verschil(p, soort, f"{a.get('jaar')} ({doc})", f"{b.get('jaar')} ({doc})", a["bedrag"], b["bedrag"]))
        per_jaar = {}
        for b in bedragen:
            per_jaar.setdefault((b.get("soort") or "", str(b.get("jaar"))), []).append(b)
        for (soort, jaar), groep in per_jaar.items():
            docs = {}
            for b in groep:
                docs.setdefault(b.get("document") or "", b)
            lijst = list(docs.values())
            for a, b in zip(lijst, lijst[1:]):
                rijen.append(_verschil(p, soort, f"{jaar}: {a.get('document')}", f"{jaar}: {b.get('document')}", a["bedrag"], b["bedrag"]))
    return rijen


def _verschil(p, soort, van, naar, a, b):
    verschil = b - a
    procent = (verschil / abs(a) * 100) if a else None
    opvallend = abs(verschil) >= DREMPEL_EURO or (procent is not None and abs(procent) >= DREMPEL_PROCENT)
    return {"post": p.get("post"), "soort": soort, "van": van, "naar": naar, "bedrag_van": a, "bedrag_naar": b,
            "verschil": verschil, "procent": round(procent, 1) if procent is not None else None, "opvallend": opvallend}


def euro(x):
    """€ 1.234.567 (Nederlandse notatie)."""
    if x is None:
        return ""
    teken = "-" if x < 0 else ""
    return f"{teken}€ {abs(round(x)):,.0f}".replace(",", ".")


def maak(c, bid):
    b = c.execute("SELECT * FROM begrotingsanalyses WHERE id = %s", (bid,)).fetchone()
    claude.schrijfplafond(c)
    ids = [int(x) for x in b["documenten"] or []]
    docs = c.execute("SELECT id, titel FROM documenten WHERE id = ANY(%s)", (ids or [0],)).fetchall()
    q = lambda sql, p=None: c.execute(sql, p).fetchall()  # noqa: E731
    vraag = f"{b['thema']} {b['vraag'] or ''}"
    passages = kennisbank.context(q, vraag, maximum=60000, documenten=len(ids) or 6, alleen=ids, fragmenten=12)
    beloften = kennisbank.context(q, vraag, maximum=12000, documenten=3,
                                  voorkeur=("coalitieakkoord", "verkiezingsprogramma-groenlinks-pvda"), fragmenten=5)
    delen = [f"# Thema\n{b['thema']}", f"# Vraag van Martijn\n{b['vraag'] or 'Wat staat er in de stukken en wat is er veranderd?'}",
             "# Gekozen documenten\n" + "\n".join(f"- {d['titel']}" for d in docs),
             "# Passages uit de gekozen documenten\n" + (passages or "Niets gevonden bij dit thema."),
             "# Passages over beloften (coalitieakkoord, programma)\n" + (beloften or "Niets gevonden."),
             "# Opdracht\n" + OPDRACHT]
    model = nieuws.instelling(c, "claude_model_schrijven", "claude-sonnet-5-5")
    profiel = nieuws.instelling(c, "nieuws_profiel", nieuws.STANDAARD_PROFIEL)
    systeem = (f"Je bent de begrotingsanalist van de fractie van PRO Helmond in de gemeenteraad van Helmond.\n\n{profiel}\n\n"
               f"Vandaag is het {whaller.lang(dt.date.today())}.")
    antwoord = claude.vraag(c, "begroting", systeem, "\n\n".join(delen), model, max_tokens=8000, timeout=400)
    inzicht = claude.json_uit(antwoord)
    inzicht = json.loads(schrijven._geen_gedachtestreepjes(json.dumps(inzicht, ensure_ascii=False)))  # pylint: disable=protected-access
    inzicht["berekend"] = bereken(inzicht.get("posten") or [])
    c.execute("""UPDATE begrotingsanalyses SET inzicht = %s, status = 'klaar', model = %s, fout = NULL,
                   titel = coalesce(nullif(%s, ''), titel), bijgewerkt = now() WHERE id = %s""",
              (json.dumps(inzicht, ensure_ascii=False), model, (inzicht.get("titel") or "")[:200], bid))


def start(pool, bid):
    def werk():
        with pool.connection() as c:
            try:
                maak(c, bid)
            except Exception as e:  # noqa: BLE001
                c.execute("UPDATE begrotingsanalyses SET status = 'fout', fout = %s, bijgewerkt = now() WHERE id = %s",
                          (f"{e}\n\n{traceback.format_exc()[-800:]}", bid))
    threading.Thread(target=werk, daemon=True).start()


def vragen_tekst(b):
    """De technische vragen uit de analyse, als tekst om een concept mee te beginnen."""
    i = b["inzicht"] or {}
    regels = [f"Begrotingshulp: {b['titel']}", ""]
    regels += [f"- {v}" for v in i.get("vragen") or []]
    if i.get("opvallend"):
        regels += ["", "Wat opviel:"] + [f"- {o}" for o in i["opvallend"]]
    return "\n".join(regels)


def word(b):
    import docx  # pylint: disable=import-outside-toplevel
    from docx.shared import Pt  # pylint: disable=import-outside-toplevel
    i = b["inzicht"] or {}
    doc = docx.Document()
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(10)
    doc.add_heading(f"Begrotingshulp: {b['titel']}", level=1)
    doc.add_paragraph(f"Thema: {b['thema']}. Gemaakt {whaller.d(b['bijgewerkt'].date())}. "
                      "Bedragen met bron uit de stukken; de verschillen zijn door Promonitor uitgerekend.")
    if i.get("toelichting"):
        doc.add_paragraph(i["toelichting"])
    if i.get("posten"):
        doc.add_heading("Bedragen", level=2)
        t = doc.add_table(rows=1, cols=5)
        t.style = "Table Grid"
        for cel, k in zip(t.rows[0].cells, ("Post", "Jaar", "Bedrag", "In het stuk", "Bron")):
            cel.text = k
        for p in i["posten"]:
            for x in p.get("bedragen") or []:
                cellen = t.add_row().cells
                for cel, w in zip(cellen, (p.get("post") or "", str(x.get("jaar") or ""), euro(_bedrag(x.get("bedrag"))),
                                          f"{x.get('bedrag_bron', '')} {x.get('eenheid_bron', '')}".strip(),
                                          f"{x.get('document', '')} {x.get('pagina', '')}".strip())):
                    cel.text = w
    if i.get("berekend"):
        doc.add_heading("Verschillen (uitgerekend)", level=2)
        t = doc.add_table(rows=1, cols=5)
        t.style = "Table Grid"
        for cel, k in zip(t.rows[0].cells, ("Post", "Van", "Naar", "Verschil", "Procent")):
            cel.text = k
        for r in i["berekend"]:
            cellen = t.add_row().cells
            for cel, w in zip(cellen, (r["post"] or "", f"{r['van']}: {euro(r['bedrag_van'])}", f"{r['naar']}: {euro(r['bedrag_naar'])}",
                                      euro(r["verschil"]), f"{r['procent']}%".replace(".", ",") if r["procent"] is not None else "")):
                cel.text = w
            if r["opvallend"]:
                for cel in cellen:
                    for run in cel.paragraphs[0].runs:
                        run.bold = True
    for kop, sleutel in (("Beloften", "beloften"), ("Wat opvalt", "opvallend"), ("Vragen voor de raad", "vragen")):
        if not i.get(sleutel):
            continue
        doc.add_heading(kop, level=2)
        for x in i[sleutel]:
            if isinstance(x, dict):
                tekst = f"{x.get('belofte')} ({x.get('bron')}). In de stukken: {x.get('in_stukken')}. {x.get('toelichting') or ''}"
            else:
                tekst = x
            doc.add_paragraph(tekst.strip(), style="List Bullet")
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue()


def bestandsnaam(b):
    basis = re.sub(r"[^\w\- ]+", "", b["titel"] or "begrotingshulp").strip().replace(" ", "-")[:60]
    return f"{dt.date.today():%Y-%m-%d}-begrotingshulp-{basis}.docx"
