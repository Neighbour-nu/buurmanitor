"""Stemadvies per raadsvergadering: per raadsvoorstel het fractiestandpunt op één A4.

Promonitor zet de hamer- en bespreekstukken van een raadsvergadering op een rij, met de woordvoerder.
Claude doet per punt een voorstel (voor, tegen, voor met kanttekening, of nog bespreken) op basis van wat
de fractie al vastlegde: fractiebesluiten, standpunten in dossiers, notities bij agendapunten, concepten en
de stukken zelf. De fractie beslist: elk voorstel is aan te passen. Daarna een Word-bestand op één A4.
"""
import datetime as dt
import io
import json
import re
import threading
import traceback

from . import agenda, claude, nieuws, notubiz, schrijven, whaller

ADVIEZEN = [("voor", "Voor"), ("voor_kanttekening", "Voor, met kanttekening"), ("tegen", "Tegen"),
            ("onthouden", "Niet deelnemen"), ("bespreken", "Nog bespreken in de fractie")]
ADVIES_NAAM = dict(ADVIEZEN)


def punten(verg, toew, raadsleden, namen):
    """De punten waarover gestemd wordt: hamer- en bespreekstukken, zonder formele punten."""
    ind = whaller.raad_indeling(verg, toew, raadsleden)
    uit = []
    for rubriek in ("bespreek", "hamer"):
        for p in ind[rubriek]:
            if agenda.is_formeel(p["onderwerp"]):
                continue
            uit.append({"punt_id": p["id"], "nummer": (p["nummer"] or "").rstrip("."), "onderwerp": p["onderwerp"], "rubriek": rubriek,
                        "woordvoerder": namen.get(p.get("woordvoerder"), p.get("woordvoerder") or ""),
                        "sub": p.get("sub", []),
                        "documenten": [{"id": d["id"], "titel": d["titel"], "url": d["url"]} for d in p.get("documenten", [])],
                        "advies": "voor" if rubriek == "hamer" else "bespreken", "motivering": "",
                        "stemverklaring": "", "bron": ""})
    return uit


def _wat_we_weten(c, onderwerp):
    """Wat de fractie al over dit onderwerp vastlegde."""
    k = whaller.sleutel(onderwerp)
    woorden = [w for w in k.split() if len(w) > 4][:6]
    pat = [f"%{w}%" for w in woorden] or [f"%{k}%"]
    regels = []
    for r in c.execute("SELECT onderwerp, notitie, gedaan FROM agenda_voortgang WHERE sleutel = %s", (k,)).fetchall():
        if r["notitie"]:
            regels.append(f"Notitie bij het agendapunt: {r['notitie']}")
    for d in c.execute("""SELECT naam, standpunt, inzicht->>'standpunt' AS afgeleid FROM dossiers
                          WHERE status = 'actief' AND (%s ILIKE '%%' || naam || '%%'
                                OR EXISTS (SELECT 1 FROM unnest(string_to_array(coalesce(zoektermen, ''), ',')) t
                                           WHERE length(trim(t)) > 3 AND %s ILIKE '%%' || trim(t) || '%%'))""",
                       (onderwerp, onderwerp)).fetchall():
        if d["standpunt"] or d["afgeleid"]:
            regels.append(f"Dossier {d['naam']}: {d['standpunt'] or d['afgeleid']}")
    for n in c.execute("SELECT datum, inhoud FROM notulen WHERE inhoud IS NOT NULL ORDER BY datum DESC LIMIT 8").fetchall():
        for p in (n["inhoud"] or {}).get("punten", []):
            tekst = " ".join([p.get("titel") or "", p.get("verslag") or ""] + (p.get("besluiten") or []))
            if sum(w in tekst.lower() for w in woorden) >= max(1, min(2, len(woorden))):
                besl = "; ".join(p.get("besluiten") or [])
                regels.append(f"Fractievergadering {n['datum']:%d-%m}: {besl or (p.get('verslag') or '')[:400]}")
        for w in (n["inhoud"] or {}).get("woordvoering", []):
            if any(x in (w.get("agendapunt") or "").lower() for x in woorden) and w.get("standpunt"):
                regels.append(f"Woordvoering {n['datum']:%d-%m}: {w.get('woordvoerder')}: {w['standpunt']}")
    for k2 in c.execute("""SELECT soort, titel, left(tekst, 1500) AS tekst FROM concepten
                           WHERE status = 'klaar' AND soort IN ('analyse', 'bijdrage', 'stemverklaring', 'amendement', 'motie')
                             AND (titel ILIKE ANY(%s) OR coalesce(bron->>'onderwerp', '') ILIKE ANY(%s))
                           ORDER BY id DESC LIMIT 3""", (pat, pat)).fetchall():
        kort = re.sub(r"\s+", " ", k2["tekst"] or "")[:1200]
        regels.append(f"Concept ({k2['soort']}) {k2['titel']}: {kort}")
    return regels


OPDRACHT = """Doe per agendapunt een voorstel voor het stemadvies van de fractie van PRO Helmond.
Regels:
- De fractie beslist. Jij doet een voorstel op basis van wat de fractie al vastlegde en van de stukken.
- Is er niets bekend over het standpunt en is het een bespreekstuk, kies dan "bespreken" en zeg wat de fractie moet afwegen.
- Hamerstukken: meestal "voor", tenzij uit de stukken of het fractiestandpunt een bezwaar blijkt.
- "voor_kanttekening" als we instemmen maar een zorg hebben die we willen uitspreken (stemverklaring).
- Motivering: één of twee korte zinnen, B2, geen gedachtestreepjes.
- Stemverklaring: alleen invullen ("ja, door ...") als een stemverklaring zinvol is; anders leeg.
- bron: waarop je het voorstel baseert (bijvoorbeeld "fractievergadering 19-10", "dossier", "alleen de stukken").

Antwoord met alleen één blok ```json:
{"punten": [{"punt_id": 123, "advies": "voor|voor_kanttekening|tegen|onthouden|bespreken",
             "motivering": "...", "stemverklaring": "", "bron": "..."}],
 "toelichting": "twee zinnen aan Martijn"}"""


def maak(c, sid):
    s = c.execute("SELECT * FROM stemadviezen WHERE id = %s", (sid,)).fetchone()
    claude.schrijfplafond(c)
    items = s["items"] or []
    delen = [f"# Raadsvergadering {whaller.lang(s['datum'])}"]
    for p in items:
        blok = [f"## Punt {p['nummer']} (punt_id {p['punt_id']}): {p['onderwerp']}",
                f"Rubriek: {'bespreekstuk' if p['rubriek'] == 'bespreek' else 'hamerstuk'}",
                f"Woordvoerder: {p['woordvoerder'] or 'geen'}"]
        if p.get("sub"):
            blok.append("Onderdelen: " + "; ".join(p["sub"]))
        weten = _wat_we_weten(c, p["onderwerp"])
        blok.append("Wat de fractie al vastlegde:\n" + ("\n".join(f"- {w}" for w in weten) if weten else "- niets gevonden"))
        if p.get("documenten"):
            tekst, _, gemist = schrijven.stukken_tekst(c, p["documenten"][:2], maximum=6000 if p["rubriek"] == "bespreek" else 2500)
            if tekst:
                blok.append("Uit de stukken:\n" + tekst)
            elif gemist:
                blok.append("De stukken konden niet worden gelezen.")
        delen.append("\n".join(blok))
    delen.append("# Opdracht\n" + OPDRACHT)
    model = nieuws.instelling(c, "claude_model_schrijven", "claude-sonnet-5-5")
    profiel = nieuws.instelling(c, "nieuws_profiel", nieuws.STANDAARD_PROFIEL)
    systeem = (f"Je bereidt het stemadvies voor van de fractie van PRO Helmond in de gemeenteraad van Helmond.\n\n{profiel}\n\n"
               f"Vandaag is het {whaller.lang(dt.date.today())}.")
    antwoord = claude.vraag(c, "stemadvies", systeem, "\n\n".join(delen), model, max_tokens=6000, timeout=400)
    uit = claude.json_uit(antwoord)
    per_punt = {str(x.get("punt_id")): x for x in uit.get("punten", [])}
    for p in items:
        x = per_punt.get(str(p["punt_id"]))
        if not x:
            continue
        if x.get("advies") in ADVIES_NAAM:
            p["advies"] = x["advies"]
        for veld in ("motivering", "stemverklaring", "bron"):
            p[veld] = schrijven._geen_gedachtestreepjes(str(x.get(veld) or ""))  # pylint: disable=protected-access
    c.execute("""UPDATE stemadviezen SET items = %s, toelichting = %s, status = 'concept', model = %s, fout = NULL,
                   bijgewerkt = now() WHERE id = %s""",
              (json.dumps(items, ensure_ascii=False), schrijven._geen_gedachtestreepjes(uit.get("toelichting", "")),  # pylint: disable=protected-access
               model, sid))


def start(pool, sid):
    def werk():
        with pool.connection() as c:
            try:
                maak(c, sid)
            except Exception as e:  # noqa: BLE001
                c.execute("UPDATE stemadviezen SET status = 'fout', fout = %s, bijgewerkt = now() WHERE id = %s",
                          (f"{e}\n\n{traceback.format_exc()[-800:]}", sid))
    threading.Thread(target=werk, daemon=True).start()


def nieuw(c, notubiz_id, personen):
    """Maakt (of vernieuwt de punten van) het stemadvies voor een raadsvergadering. Geeft het id terug."""
    verg = notubiz.vergadering(notubiz_id)
    toew = c.execute("SELECT * FROM toewijzingen").fetchall()
    raadsleden = {p["whaller"] for p in personen if (p["rol"] or "").startswith("raadslid")}
    namen = {p["whaller"]: p["naam"] for p in personen}
    items = punten(verg, toew, raadsleden, namen)
    oud = c.execute("SELECT id, items FROM stemadviezen WHERE notubiz_id = %s", (notubiz_id,)).fetchone()
    if oud:
        # Bewaar wat al was ingevuld voor punten die er nog steeds zijn.
        bekend = {str(p["punt_id"]): p for p in oud["items"] or []}
        for p in items:
            o = bekend.get(str(p["punt_id"]))
            if o:
                p.update({k: o[k] for k in ("advies", "motivering", "stemverklaring", "bron") if k in o})
        c.execute("UPDATE stemadviezen SET items = %s, bijgewerkt = now() WHERE id = %s",
                  (json.dumps(items, ensure_ascii=False), oud["id"]))
        return oud["id"]
    r = c.execute("""INSERT INTO stemadviezen (notubiz_id, datum, titel, url, items) VALUES (%s, %s, %s, %s, %s)
                     RETURNING id""",
                  (notubiz_id, verg["start"].date(), verg["naam"], verg["url"], json.dumps(items, ensure_ascii=False))).fetchone()
    return r["id"]


def word(s):
    """Stemadvies op één A4: compacte tabel, liggend als er veel punten zijn."""
    import docx  # pylint: disable=import-outside-toplevel
    from docx.enum.section import WD_ORIENT  # pylint: disable=import-outside-toplevel
    from docx.shared import Cm, Pt, RGBColor  # pylint: disable=import-outside-toplevel
    items = s["items"] or []
    doc = docx.Document()
    sectie = doc.sections[0]
    for kant in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sectie, kant, Cm(1.5))
    if len(items) > 14:
        sectie.orientation = WD_ORIENT.LANDSCAPE
        sectie.page_width, sectie.page_height = sectie.page_height, sectie.page_width
    stijl = doc.styles["Normal"]
    stijl.font.name, stijl.font.size = "Calibri", Pt(9 if len(items) > 10 else 10)
    stijl.paragraph_format.space_after = Pt(0)
    kop = doc.add_paragraph()
    r = kop.add_run(f"Stemadvies PRO Helmond: {s['titel'] or 'Raadsvergadering'} {whaller.d(s['datum'])}")
    r.bold, r.font.size = True, Pt(14)
    status = "Vastgesteld door de fractie" if s["status"] == "vastgesteld" else "Concept, nog niet vastgesteld"
    doc.add_paragraph(status).runs[0].italic = True
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    breedtes = [Cm(1.0), Cm(7.0), Cm(3.0), Cm(5.5), Cm(2.5)]
    for cel, tekst, b in zip(t.rows[0].cells, ("Nr", "Agendapunt", "Advies", "Toelichting", "Stemverklaring"), breedtes):
        cel.text, cel.width = tekst, b
        cel.paragraphs[0].runs[0].bold = True
    for rubriek, label in (("bespreek", "Bespreekstukken"), ("hamer", "Hamerstukken")):
        groep = [p for p in items if p["rubriek"] == rubriek]
        if not groep:
            continue
        rij = t.add_row().cells
        samen = rij[0].merge(rij[4])
        samen.text = label
        samen.paragraphs[0].runs[0].bold = True
        for p in groep:
            cellen = t.add_row().cells
            waarden = (p["nummer"] or "", p["onderwerp"] + (f" ({p['woordvoerder']})" if p.get("woordvoerder") else ""),
                       ADVIES_NAAM.get(p["advies"], p["advies"]), p.get("motivering") or "", p.get("stemverklaring") or "")
            for cel, w, b in zip(cellen, waarden, breedtes):
                cel.text, cel.width = w, b
            if p["advies"] in ("tegen", "bespreken"):
                for run in cellen[2].paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xD0, 0x00, 0x00)
    uit = io.BytesIO()
    doc.save(uit)
    return uit.getvalue()


def bestandsnaam(s):
    return f"{s['datum']:%Y-%m-%d}-stemadvies-raadsvergadering{'' if s['status'] == 'vastgesteld' else '-concept'}.docx"
