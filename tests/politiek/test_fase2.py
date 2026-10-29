"""Tests voor begrotingshulp, dashboard, planner, stemadvies en feedback (zonder netwerk en zonder Claude).

Draaien: python -m tests.test_fase2
"""
import datetime as dt
import io

from app.politiek import begroting, dashboard, schrijven, stemadvies, taken

TZ = taken.TZ


def test_begroting_verschillen():
    posten = [{"post": "Sport", "bedragen": [
        {"jaar": "2028", "bedrag": 7550000, "soort": "lasten", "document": "Begroting"},
        {"jaar": "2027", "bedrag": 7290000, "soort": "lasten", "document": "Begroting"},
        {"jaar": "2027", "bedrag": "7.000.000", "soort": "lasten", "document": "Kalender"}]}]
    rijen = begroting.bereken(posten)
    jaren = next(r for r in rijen if r["van"].startswith("2027 (Begroting"))
    assert jaren["verschil"] == 260000 and jaren["procent"] == 3.6 and jaren["opvallend"]
    docs = next(r for r in rijen if r["van"].startswith("2027: "))
    assert docs["verschil"] == -290000


def test_euro():
    assert begroting.euro(1234567) == "€ 1.234.567"
    assert begroting.euro(-5000) == "-€ 5.000"
    assert dashboard.euro(1.5) == "€ 1,50"


def test_eerste_werkdag():
    assert dashboard.eerste_werkdag(dt.date(2026, 11, 17)) == dt.date(2026, 11, 2)  # 1 november is een zondag
    assert dashboard.eerste_werkdag(dt.date(2026, 12, 5)) == dt.date(2026, 12, 1)


def test_zondagavond():
    zondag = dt.datetime(2026, 11, 1, 20, 30, tzinfo=TZ)
    assert taken._zondagavond(None, zondag) == (dt.time(20, 0), 24)
    assert taken._zondagavond(None, zondag - dt.timedelta(days=1)) is False
    assert taken.is_aan_de_beurt(None, None, None, zondag, dt.time(20, 0), 24)
    assert not taken.is_aan_de_beurt(zondag - dt.timedelta(minutes=5), None, None, zondag, dt.time(20, 0), 24)
    # De gewone briefing stopt nog steeds om twaalf uur.
    assert not taken.is_aan_de_beurt(None, None, None, zondag.replace(hour=13), dt.time(8, 0))


def test_stemadvies_word_een_pagina_breed():
    items = [{"punt_id": i, "nummer": str(i), "onderwerp": f"Punt {i}", "rubriek": "hamer" if i > 3 else "bespreek",
              "woordvoerder": "", "advies": "voor", "motivering": "", "stemverklaring": ""} for i in range(1, 6)]
    inhoud = stemadvies.word({"items": items, "titel": "Raadsvergadering", "datum": dt.date(2026, 10, 27), "status": "concept"})
    assert inhoud[:2] == b"PK"


def test_feedback_opmerkingen_in_kantlijn():
    import docx
    d = docx.Document()
    d.add_paragraph("Voorzitter,")
    d.add_paragraph("Wij steunen de fusie van harte.")
    buf = io.BytesIO()
    d.save(buf)
    k = {"invoer": {}, "meta": {"opmerkingen": [{"citaat": "steunen de fusie", "opmerking": "Noem de waarde — eerst.", "voorstel": ""}]}}
    inhoud, n = schrijven.word_met_opmerkingen(k, buf.getvalue())
    assert n == 1
    terug = docx.Document(io.BytesIO(inhoud))
    opm = list(terug.comments)
    assert len(opm) == 1 and "—" not in opm[0].text


if __name__ == "__main__":
    for naam, functie in list(globals().items()):
        if naam.startswith("test_"):
            functie()
            print("ok", naam)
